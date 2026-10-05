"""Offline checks written as ordinary functions. No live Gemini requests."""
import io
import json
import uuid
from pathlib import Path
from unittest.mock import Mock
import pytest
from pharma import agent, retrieval, provider, documents, evaluation, web
from pharma.schema import ANSWER, CLAIM, PLAN, REVIEW, RESPONSE, validate_data
from pharma.domain import DISCLAIMER

ROOT = Path(__file__).resolve().parents[1]
DOCS = retrieval.load_documents(ROOT / 'knowledge/metformin.json')
WORK = ROOT / '.test-work'


def new_folder(name):
    path = WORK / (name + '-' + uuid.uuid4().hex)
    path.mkdir(parents=True, exist_ok=True)
    return path


def test_embedding(gemini, text, task):
    return [0., 1.] if 'topiramate' in text.lower() else [1., 0.]

# This is an embedding fixture function, not a test itself.
test_embedding.__test__ = False


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.setenv('GEMINI_API_KEY', '')
    monkeypatch.setattr(retrieval, 'embed_text', test_embedding)


@pytest.fixture(scope='module')
def database_path():
    return new_folder('functions-chroma')


@pytest.fixture
def data(database_path):
    gemini = {'embedding_model': 'controlled-test-embedding'}
    store = retrieval.open_store(database_path)
    retrieval.index_documents(store, DOCS, gemini)
    yield gemini, store
    retrieval.close_store(store)


@pytest.fixture
def client():
    return web.app.test_client()


def plan():
    return {'safe': True, 'in_domain': True, 'search_query': 'metformin study'}


def claim_answer(quote=None):
    return {'topic': 'metformin', 'status': 'answered',
            'summary': [{'text': 'The study lasted 29 weeks.', 'evidence': [
                {'source_id': DOCS[0]['source_id'], 'quote': quote or DOCS[0]['text']}]}],
            'risks': [], 'interactions': [], 'explanation': 'Based on retrieved label.'}


def generate_sequence(monkeypatch, outputs):
    def generate(gemini, system, prompt, schema):
        return validate_data(outputs.pop(0), schema)
    monkeypatch.setattr(agent, 'generate_answer', generate)


def test_persistent_cosine_search(data):
    gemini, store = data
    retrieval.check_index(store, DOCS, gemini)
    hits = retrieval.search(store, 'topiramate', gemini, min_score=.9)
    assert [h['source_id'] for h in hits] == ['metformin-interaction']
    reopened = retrieval.open_store(store['path'])
    retrieval.check_index(reopened, DOCS, gemini)
    assert retrieval.search(reopened, 'study', gemini)


def test_grounded_answer(data, monkeypatch):
    generate_sequence(monkeypatch, [plan(), claim_answer(), {'supported': True, 'safe': True, 'sufficient': True}])
    result = agent.ask_question('Summarize metformin research', *data)
    assert result['status'] == 'answered'
    assert result['disclaimer'] == DISCLAIMER
    assert result['confidence'] == 'limited_evidence'
    assert result['sources'][0]['source_id'] == 'metformin-study'
    validate_data(result, RESPONSE)


def test_personal_advice_before_model(data, monkeypatch):
    generate = Mock(side_effect=AssertionError('No model call expected'))
    monkeypatch.setattr(agent, 'generate_answer', generate)
    assert agent.ask_question('Should I stop taking metformin?', *data)['status'] == 'refused'
    generate.assert_not_called()


def test_model_safety(data, monkeypatch):
    generate_sequence(monkeypatch, [{'safe': False, 'in_domain': True, 'search_query': ''}])
    assert agent.ask_question('Choose treatment for this patient', *data)['status'] == 'refused'


def test_empty_retrieval(data, monkeypatch):
    generate_sequence(monkeypatch, [plan()])
    monkeypatch.setattr(agent, 'search', lambda *args, **kwargs: [])
    assert agent.ask_question('Unknown drug', *data)['status'] == 'insufficient_evidence'


def test_unknown_citation():
    answer = claim_answer()
    answer['summary'][0]['evidence'][0]['source_id'] = 'invented'
    assert not agent.validate_grounding(answer, DOCS)


def test_fabricated_quote(data, monkeypatch):
    generate_sequence(monkeypatch, [plan(), claim_answer('Invented mortality benefit')])
    answer = agent.ask_question('Summarize metformin', *data)
    assert answer['status'] == 'insufficient_evidence'
    assert not answer['summary']


def test_semantic_review(data, monkeypatch):
    generate_sequence(monkeypatch, [plan(), claim_answer(), {'supported': False, 'safe': True, 'sufficient': True}])
    assert agent.ask_question('Summarize metformin', *data)['status'] == 'insufficient_evidence'


def test_missing_pair(data, monkeypatch):
    answer = {**claim_answer(), 'status': 'insufficient_evidence', 'summary': []}
    generate_sequence(monkeypatch, [plan(), answer])
    hits = []
    result = agent.ask_question('Does metformin interact with QX-997?', *data, retrieved=hits)
    assert result['status'] == 'insufficient_evidence'
    assert hits


def test_abstention_has_no_claims():
    answer = claim_answer()
    answer['status'] = 'insufficient_evidence'
    assert not agent.validate_grounding(answer, DOCS)


def test_schema_rejects_bad_claims():
    with pytest.raises(ValueError):
        validate_data({'text': 'No evidence', 'evidence': []}, CLAIM)
    with pytest.raises(ValueError):
        validate_data({**plan(), 'unexpected': True}, PLAN)


def test_stale_index(data):
    gemini, store = data
    changed = json.loads(json.dumps(DOCS))
    changed[0]['text'] += 'changed'
    with pytest.raises(ValueError):
        retrieval.check_index(store, changed, gemini)


def test_bad_vectors():
    for vector in ([], [0, 0], [float('nan'), 1]):
        with pytest.raises(ValueError):
            retrieval.normalize(vector)


def test_unsafe_output():
    answer = claim_answer()
    answer['summary'][0]['text'] = 'You should stop taking metformin.'
    assert not agent.validate_grounding(answer, DOCS)


def test_all_drugs(data, monkeypatch):
    gemini, store = data
    drugs = ['metformin', 'glipizide', 'amlodipine', 'lisinopril', 'alendronate']
    monkeypatch.setattr(retrieval, 'embed_text', lambda gemini, text, task: [float(d in text.lower()) for d in drugs])
    docs = retrieval.load_documents(ROOT / 'knowledge')
    retrieval.index_documents(store, docs, gemini)
    for drug in drugs:
        hits = retrieval.search(store, drug, gemini, min_score=.9)
        assert hits and all(h['drug'].lower() == drug for h in hits)
    assert store['collection'].count() == 15
    assert (store['path'] / 'chroma.sqlite3').is_file()


def test_corpus_coverage():
    docs = retrieval.load_documents(ROOT / 'knowledge')
    assert len(docs) == 15
    assert {d['drug'] for d in docs} == {'Metformin', 'Glipizide', 'Amlodipine', 'Lisinopril', 'Alendronate'}
    assert {d['category'] for d in docs} == {'diabetes', 'hypertension', 'bone-health'}
    assert len({d['url'] for d in docs}) == 6


def test_real_paper():
    paper = [d for d in retrieval.load_documents(ROOT / 'knowledge') if d['kind'] == 'research-paper']
    assert len(paper) == 2
    assert all('11832527' in d['url'] and '10.1056/NEJMoa012512' in d['section'] for d in paper)
    text = ' '.join(d['text'] for d in paper)
    assert all(fact in text for fact in ['3234', '2.8', '31%', '58%', 'without diabetes'])


def test_evaluation_cases():
    cases = json.loads((ROOT / 'evaluation/cases.json').read_text())
    known = {d['source_id'] for d in retrieval.load_documents(ROOT / 'knowledge')}
    assert len(cases) == 16
    assert all(set(c['expected_sources']) <= known for c in cases)


def test_sdk_json_schema():
    from google.genai import types
    config = types.GenerateContentConfig(response_mime_type='application/json', response_json_schema=ANSWER)
    assert config.response_json_schema == ANSWER


def pdf_bytes(text=True):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if text:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 20 250 Td (Research document extraction test.) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_txt_upload():
    assert documents.clean_filename('../../paper.txt') == 'paper.txt'
    assert documents.clean_filename('C:\\unsafe\\paper.txt') == 'paper.txt'
    document_id, records, warnings = documents.make_records('paper.txt', 'Research عربي'.encode())
    assert len(document_id) == 64 and 'عربي' in records[0]['text']
    assert records[0]['title'] == 'paper.txt' and not warnings


def test_invalid_documents():
    for name, content in [('x.doc', b'a'), ('x.exe', b'a'), ('x.txt', b''), ('x.txt', b'\xff'), ('x.txt', b'a\x00b')]:
        with pytest.raises(ValueError):
            documents.make_records(name, content)


def test_pdf_page():
    _, records, _ = documents.make_records('research.pdf', pdf_bytes())
    assert 'Research document extraction test.' in records[0]['text']
    assert 'Page 1' in records[0]['section']


def test_no_ocr():
    with pytest.raises(ValueError, match='No extractable text'):
        documents.make_records('scan.pdf', pdf_bytes(False))


def test_docx():
    from docx import Document
    doc = Document()
    doc.add_paragraph('Research paragraph.')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Group'
    table.cell(0, 1).text = 'Outcome'
    buffer = io.BytesIO()
    doc.save(buffer)
    _, records, _ = documents.make_records('research.docx', buffer.getvalue())
    text = ' '.join(r['text'] for r in records)
    assert all(term in text for term in ('Research paragraph.', 'Group', 'Outcome'))
    assert 'Document text' in records[0]['section']


def test_chunks():
    _, records, _ = documents.make_records('long.txt', b'research evidence ' * 300)
    assert len(records) > 1 and all(len(r['text']) <= 700 for r in records)
    with pytest.raises(ValueError):
        documents.make_records('large.txt', b'a' * 200001)


def test_upload_persistence():
    folder = new_folder('upload')
    gemini = {'embedding_model': 'upload-test-model'}
    store = retrieval.open_store(folder / 'chroma')
    info = documents.upload_document(folder / 'uploads', 'paper.txt', b'Research evidence.', gemini, store)
    reopened = retrieval.open_store(folder / 'chroma')
    retrieval.check_index(reopened, DOCS, gemini)
    hits = retrieval.search(reopened, 'research', gemini, document_id=info['id'])
    assert len(hits) == 1 and hits[0]['title'] == 'paper.txt'
    assert '/api/documents/' in hits[0]['url']
    documents.upload_document(folder / 'uploads', 'paper.txt', b'Research evidence.', gemini, reopened)
    assert reopened['uploads'].count() == 1
    assert len(documents.list_uploads(folder / 'uploads')) == 1
    assert not retrieval.search(reopened, 'research', gemini, document_id='0' * 64)


def test_failed_upload(monkeypatch):
    folder = new_folder('failed-upload')
    monkeypatch.setattr(documents, 'add_documents', Mock(side_effect=RuntimeError('unavailable')))
    with pytest.raises(RuntimeError):
        documents.upload_document(folder, 'paper.txt', b'Extractable text', {'embedding_model': 'test'}, {})
    assert documents.list_uploads(folder) == []


def test_static_routes(client):
    for path, content in [('/', b'Research question'), ('/app.js', b'/api/ask'), ('/style.css', b'@media')]:
        result = client.get(path)
        assert result.status_code == 200 and content in result.data
        assert "script-src 'self'" in result.headers['Content-Security-Policy']


def test_status_secret(client, monkeypatch):
    monkeypatch.setenv('GEMINI_API_KEY', 'private-test-secret')
    result = client.get('/api/status')
    assert result.status_code == 200 and b'private-test-secret' not in result.data
    assert result.json['key_configured']


def test_web_invalid_input(client):
    for payload in [b'bad-json', b'{"query":""}', b'[]', b'{"query":8}']:
        assert client.post('/api/ask', data=payload, content_type='application/json').status_code == 400
    assert client.post('/api/ask', json={'query': 'research'}, headers={'Origin': 'https://untrusted.example'}).status_code == 403


def test_web_missing_key(client):
    result = client.post('/api/ask', json={'query': 'metformin'})
    assert result.status_code == 503 and 'GEMINI_API_KEY' in result.json['error']


def test_web_pipeline(client, monkeypatch):
    monkeypatch.setenv('GEMINI_API_KEY', 'test')
    gemini, store = {'embedding_model': 'test'}, {'collection': None}
    monkeypatch.setattr(web, 'connect_gemini', lambda: gemini)
    monkeypatch.setattr(web, 'open_store', lambda path: store)
    monkeypatch.setattr(web, 'check_index', lambda *args: None)
    ask = Mock(return_value=agent.abstain('unknown', 'Insufficient evidence'))
    close = Mock()
    monkeypatch.setattr(web, 'ask_question', ask)
    monkeypatch.setattr(web, 'close_store', close)
    result = client.post('/api/ask', json={'query': 'unknown'})
    assert result.status_code == 200 and result.json['status'] == 'insufficient_evidence'
    ask.assert_called_once_with('unknown', gemini, store, 4, .35, document_id=None)
    close.assert_called_once_with(store)


def test_web_upload_validation(client):
    result = client.post('/api/upload', data=b'Pharma document text', content_type='application/octet-stream', headers={'X-Filename': 'paper.txt'})
    assert result.status_code == 503 and 'GEMINI_API_KEY' in result.json['error']
    result = client.post('/api/upload', data=b'unsupported', content_type='application/octet-stream', headers={'X-Filename': 'paper.exe'})
    assert result.status_code == 400


def test_web_upload_download(client, monkeypatch):
    folder = new_folder('http-upload')
    monkeypatch.setattr(web, 'UPLOAD_ROOT', folder)
    monkeypatch.setenv('GEMINI_API_KEY', 'test')
    monkeypatch.setattr(web, 'connect_gemini', lambda: {'embedding_model': 'http-test-model'})
    monkeypatch.setattr(web, 'open_store', lambda path: {'collection': None})
    monkeypatch.setattr(web, 'close_store', lambda store: None)
    monkeypatch.setattr(documents, 'add_documents', lambda *args: None)
    result = client.post('/api/upload', data=b'Research content', content_type='application/octet-stream', headers={'X-Filename': 'paper.txt'})
    assert result.status_code == 200
    document_id = result.json['document']['id']
    assert len(client.get('/api/documents').json['documents']) == 1
    result = client.get('/api/documents/' + document_id)
    assert result.status_code == 200 and result.data == b'Research content'
    assert 'attachment' in result.headers['Content-Disposition']


def test_generation_validation(monkeypatch):
    client = Mock()
    client.models.generate_content.return_value.text = json.dumps(plan())
    gemini = {'client': client, 'model': 'test'}
    assert provider.generate_answer(gemini, 'system', 'prompt', PLAN) == plan()
    config = client.models.generate_content.call_args.kwargs['config']
    assert config.response_json_schema == PLAN
    client.models.generate_content.return_value.text = '{"safe":true}'
    with pytest.raises(ValueError):
        provider.generate_answer(gemini, 'system', 'prompt', PLAN)


def test_evaluation_functions(monkeypatch):
    answer = agent.abstain('unknown', 'Insufficient evidence')
    monkeypatch.setattr(evaluation, 'ask_question', lambda *args, **kwargs: answer)
    case = {'id': 'missing', 'query': 'unknown', 'status': 'insufficient_evidence', 'expected_sources': [], 'field': 'summary', 'required_terms': []}
    report = evaluation.evaluate({}, {}, [case])
    assert report['passed'] == report['total'] == 1


def test_cli_missing_key(monkeypatch, capsys):
    from pharma.__main__ import main
    monkeypatch.setattr('sys.argv', ['pharma', 'index'])
    assert main() == 1
    assert 'GEMINI_API_KEY' in capsys.readouterr().err


def test_web_scoped_question(client, monkeypatch):
    monkeypatch.setenv('GEMINI_API_KEY', 'test')
    monkeypatch.setattr(web, 'connect_gemini', lambda: {})
    monkeypatch.setattr(web, 'open_store', lambda path: {'collection': None})
    monkeypatch.setattr(web, 'check_index', lambda *args: None)
    monkeypatch.setattr(web, 'close_store', lambda *args: None)
    ask = Mock(return_value=agent.abstain('paper', 'Missing evidence'))
    monkeypatch.setattr(web, 'ask_question', ask)
    result = client.post('/api/ask', json={'query': 'paper', 'document_id': 'a' * 64})
    assert result.status_code == 200
    assert ask.call_args.kwargs['document_id'] == 'a' * 64


def test_no_custom_classes():
    import ast
    for path in (ROOT / 'pharma').glob('*.py'):
        tree = ast.parse(path.read_text(encoding='utf-8'))
        assert not any(isinstance(node, ast.ClassDef) for node in ast.walk(tree)), path.name


def test_safe_error_messages():
    from pharma.errors import error_message
    for code, phrase in [(403, 'API access'), (429, 'quota'), (404, 'model not found'), (400, 'rejected'), (503, 'temporarily')]:
        exc = RuntimeError('private-key-must-not-be-shown')
        exc.code = code
        message = error_message(exc)
        assert phrase in message and 'private-key' not in message
    assert 'private-key' not in error_message(RuntimeError('private-key'))


def test_web_missing_index_message(client, monkeypatch):
    monkeypatch.setenv('GEMINI_API_KEY', 'test')
    monkeypatch.setattr(web, 'connect_gemini', lambda: {})
    monkeypatch.setattr(web, 'open_store', lambda path: {'collection': None})
    monkeypatch.setattr(web, 'close_store', lambda *args: None)
    monkeypatch.setattr(web, 'check_index', Mock(side_effect=ValueError('Missing vector index. Run python -m pharma index.')))
    result = client.post('/api/ask', json={'query': 'metformin'})
    assert result.status_code == 503
    assert result.json['error'] == 'Missing vector index. Run python -m pharma index.'


def test_delete_document_all_models(data):
    gemini, store = data
    folder = new_folder('delete-upload')
    first = documents.upload_document(folder, 'remove.txt', b'Remove this evidence.', gemini, store)
    second = documents.upload_document(folder, 'keep.txt', b'Keep this evidence.', gemini, store)
    other_model = {'embedding_model': 'another-upload-model'}
    documents.upload_document(folder, 'remove.txt', b'Remove this evidence.', other_model, store)
    documents.delete_document(folder, first['id'], store)
    assert not (folder / first['id']).exists()
    assert (folder / second['id'] / 'original.txt').exists()
    assert store['collection'].count() == len(DOCS)
    for collection in store['client'].list_collections():
        if collection.name.startswith('uploads-'):
            assert not collection.get(where={'document_id': first['id']})['ids']


def test_delete_invalid_path():
    folder = new_folder('invalid-delete')
    marker = folder / 'keep.txt'
    marker.write_text('keep', encoding='utf-8')
    with pytest.raises(ValueError):
        documents.delete_document(folder, '../keep.txt', {})
    with pytest.raises(FileNotFoundError):
        documents.delete_document(folder, 'a' * 64, {})
    assert marker.read_text(encoding='utf-8') == 'keep'


def test_web_delete_document(client, monkeypatch):
    folder = new_folder('http-delete')
    store = retrieval.open_store(folder / 'chroma')
    gemini = {'embedding_model': 'http-delete-model'}
    info = documents.upload_document(folder / 'uploads', 'remove.txt', b'Remove this evidence.', gemini, store)
    monkeypatch.setattr(web, 'UPLOAD_ROOT', folder / 'uploads')
    monkeypatch.setattr(web, 'open_store', lambda path: store)
    path = '/api/documents/' + info['id'] + '/delete'
    assert client.post(path, headers={'Origin': 'https://untrusted.example'}).status_code == 403
    assert client.post(path).status_code == 200
    assert client.get('/api/documents').json['documents'] == []
    assert client.get('/api/documents/' + info['id']).status_code == 404
    assert client.post(path).status_code == 404
